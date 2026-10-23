"""Sample files for Hearth Office's tests, made by independent libraries (python-docx, openpyxl, python-pptx),
so the tests check that Hearth reads what other programs write, and that what Hearth writes reads back in them."""
import io
import os
import struct
import zlib


def png(w=40, h=30, rgb=(255, 138, 61)):
    """A small solid-colour PNG."""
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))

    def chunk(t, d):
        c = struct.pack('>I', len(d)) + t + d
        return c + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')


def docx(path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = Mm(20)
    s.top_margin = s.bottom_margin = Mm(22)
    s.header.paragraphs[0].text = 'Hearth header text'
    s.footer.paragraphs[0].text = 'Hearth footer text'
    d.add_heading('Quarterly report', level=1)
    p = d.add_paragraph('Plain start, ')
    r = p.add_run('bold words'); r.bold = True
    p.add_run(', ')
    r = p.add_run('italic words'); r.italic = True
    p.add_run(', ')
    r = p.add_run('red text'); r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run(', ')
    r = p.add_run('big Georgia'); r.font.size = Pt(20); r.font.name = 'Georgia'
    p.add_run(' and ')
    r = p.add_run('underlined'); r.underline = True
    p.add_run('.')
    c = d.add_paragraph('A centred line')
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_heading('Second level heading', level=2)
    d.add_paragraph('First bullet', style='List Bullet')
    d.add_paragraph('Second bullet', style='List Bullet')
    d.add_paragraph('Step one', style='List Number')
    d.add_paragraph('Step two', style='List Number')
    t = d.add_table(rows=3, cols=3)
    t.style = 'Table Grid'
    for i in range(3):
        for j in range(3):
            t.cell(i, j).text = f'R{i + 1}C{j + 1}'
    d.add_paragraph('After the table, a picture:')
    d.add_picture(io.BytesIO(png()), width=Mm(30))
    # A hyperlink (python-docx has no helper: written as XML).
    hp = d.add_paragraph('Visit ')
    rel = d.part.relate_to('https://hearthos.fun', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    rr = OxmlElement('w:r'); t_ = OxmlElement('w:t'); t_.text = 'the Hearth site'; rr.append(t_); h.append(rr)
    hp._p.append(h)
    d.add_paragraph('Last paragraph with the word Ünïcödé and an em dash — here.')
    d.save(path)


def xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = 'Budget'
    ws.append(['Item', 'Cost', 'Qty', 'Total'])
    rows = [('Bread', 2.5, 4), ('Milk', 1.2, 6), ('Cheese', 7.25, 1), ('Apples', 0.4, 12)]
    for i, (a, b, c) in enumerate(rows, start=2):
        ws.append([a, b, c, f'=B{i}*C{i}'])
    ws['C7'] = 'Sum'
    ws['D7'] = '=SUM(D2:D5)'
    ws['C8'] = 'Avg'
    ws['D8'] = '=ROUND(AVERAGE(B2:B5),2)'
    ws['C9'] = 'Look'
    ws['D9'] = '=VLOOKUP("Milk",A2:B5,2,FALSE)'
    ws['C10'] = 'If'
    ws['D10'] = '=IF(D7>20,"over","under")'
    ws['A1'].font = Font(bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor='2F5496')
    ws['B1'].font = Font(bold=True, italic=True)
    ws['D7'].number_format = '0.00'
    ws['B2'].number_format = '"$"#,##0.00'
    ws['A3'].alignment = Alignment(horizontal='right')
    thin = Side(style='thin', color='000000')
    ws['A5'].border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.column_dimensions['A'].width = 18
    ws.freeze_panes = 'A2'
    ws.merge_cells('F1:G1')
    ws['F1'] = 'Merged'
    ws2 = wb.create_sheet('Notes')
    ws2['A1'] = 'Second sheet'
    ws2['A2'] = '=Budget!D7*2'
    wb.save(path)


def csv(path):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        f.write('Name,Score,City\n"Smith, Jo",91,Paris\nAli,78.5,"New\nYork"\nZoë,64,Kraków\n')


def pptx(path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = 'Hearth Slides test'
    s.placeholders[1].text = 'A subtitle line'
    s.notes_slide.notes_text_frame.text = 'Speaker notes for slide one'
    s2 = p.slides.add_slide(p.slide_layouts[1])
    s2.shapes.title.text = 'Bullet slide'
    body = s2.placeholders[1].text_frame
    body.text = 'First point'
    for t in ('Second point', 'Third point'):
        body.add_paragraph().text = t
    s3 = p.slides.add_slide(p.slide_layouts[6])
    tb = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1.2))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = 'Big red words'
    run.font.size = Pt(40); run.font.bold = True; run.font.color.rgb = RGBColor(0xC0, 0, 0)
    shp = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(3), Inches(3), Inches(2))
    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(0x1F, 0x9A, 0x55)
    shp.text_frame.text = 'In a box'
    s3.shapes.add_picture(io.BytesIO(png(80, 60)), Inches(6), Inches(3), Inches(2))
    p.save(path)


def make_all(folder):
    os.makedirs(folder, exist_ok=True)
    out = {}
    for name, fn in (('sample.docx', docx), ('sample.xlsx', xlsx), ('sample.csv', csv), ('sample.pptx', pptx)):
        fp = os.path.join(folder, name)
        fn(fp)
        out[name] = fp
    return out


if __name__ == '__main__':
    import sys
    print(make_all(sys.argv[1] if len(sys.argv) > 1 else 'samples'))
