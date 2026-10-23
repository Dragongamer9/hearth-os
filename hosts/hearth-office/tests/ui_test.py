"""Hearth Office's page tests: the real page (ui/) in a browser engine (Chromium and WebKit through Playwright),
with the stand-in for the Rust side (core.js, H.mock) keeping files in memory.

Each app opens sample files made by other programs (samples.py: python-docx, openpyxl, python-pptx), checks what it
shows, saves them in every format it writes, and the files are read back with those same libraries and compared.
Screenshots go to --shots.

  python3 tests/ui_test.py [--browser chromium|webkit] [--shots DIR] [--only write,sheets,slides]
"""
import argparse
import base64
import functools
import http.server
import io
import json
import os
import sys
import tempfile
import threading
import traceback
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
UI = os.path.join(os.path.dirname(HERE), 'ui')
sys.path.insert(0, HERE)
import samples  # noqa: E402

FAILS = []
PASSES = []


def check(cond, what):
    (PASSES if cond else FAILS).append(what)
    print(('ok   ' if cond else 'FAIL ') + what, flush=True)
    return cond


class Handler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def translate_path(self, path):
        p = super().translate_path(path)
        # Local runs may not have SheetJS vendored yet: HEARTH_TEST_XLSX points at a copy.
        if p.endswith('vendor/xlsx.full.min.js') and not os.path.exists(p) and os.environ.get('HEARTH_TEST_XLSX'):
            return os.environ['HEARTH_TEST_XLSX']
        return p


def serve():
    srv = http.server.ThreadingHTTPServer(('127.0.0.1', 0), functools.partial(Handler, directory=UI))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


class App:
    def __init__(self, browser, base, app, shots):
        self.page = browser.new_page(viewport={'width': 1280, 'height': 820})
        self.errors = []
        self.page.on('pageerror', lambda e: self.errors.append(str(e)))
        self.page.on('console', lambda m: m.type == 'error' and self.errors.append(m.text))
        self.page.goto(f'{base}/index.html?app={app}')
        self.page.wait_for_function('window.H && H.app && H.app.ready && document.querySelector("#start:not([hidden])")', timeout=30000)
        self.page.evaluate('H.app.ready')
        self.shots = shots
        self.name = app

    def shot(self, name):
        if self.shots:
            self.page.screenshot(path=os.path.join(self.shots, f'{self.name}-{name}.png'))

    def put(self, path, data):
        self.page.evaluate('([p, d]) => { H.mock.files[p] = d; }', [path, base64.b64encode(data).decode()])

    def get(self, path):
        d = self.page.evaluate('(p) => H.mock.files[p] || null', path)
        return base64.b64decode(d) if d else None

    def open(self, path, data):
        self.put(path, data)
        return self.page.evaluate('(p) => H.openPath(p, { here: true, throw: true }).then(() => true, (e) => String(e))', path)

    def save_as(self, path):
        ext = path.rsplit('.', 1)[-1]
        return self.page.evaluate('async ([p, ext]) => { const b = await H.app.save(ext); H.mock.files[p] = H.bytesToB64(b); return b.length; }', [path, ext])

    def js(self, code, arg=None):
        return self.page.evaluate(code, arg)


# ------------------------------------------------------------------ Write
def test_write(browser, base, files, shots, tmp):
    from docx import Document
    a = App(browser, base, 'write', shots)
    a.shot('start')
    r = a.open('/docs/sample.docx', open(files['sample.docx'], 'rb').read())
    check(r is True, f'write: opens a .docx made by python-docx ({r})')
    html = a.js('W => H.app.ed.innerHTML', None)
    for want in ['Quarterly report', 'bold words', 'italic words', 'R2C3', 'Step two', 'Ünïcödé']:
        check(want in html, f'write: the .docx shows "{want}"')
    check(a.js('() => !!H.app.ed.querySelector("h1")'), 'write: Heading 1 became a heading')
    check(a.js('() => !!H.app.ed.querySelector("h2")'), 'write: Heading 2 became a heading')
    check(a.js('() => !!H.app.ed.querySelector("b")'), 'write: bold kept')
    check(a.js('() => !!H.app.ed.querySelector("i")'), 'write: italic kept')
    check(a.js('() => !!H.app.ed.querySelector("u")'), 'write: underline kept')
    check(a.js('() => [...H.app.ed.querySelectorAll("span")].some(s => /192, 0, 0|c00000/i.test(s.style.color))'), 'write: red colour kept')
    check(a.js('() => [...H.app.ed.querySelectorAll("span")].some(s => s.style.fontSize === "20pt" && /Georgia/.test(s.style.fontFamily))'), 'write: size and font kept')
    check(a.js('() => [...H.app.ed.querySelectorAll("p")].some(p => p.style.textAlign === "center" && /centred/.test(p.textContent))'), 'write: centred paragraph kept')
    check(a.js('() => H.app.ed.querySelectorAll("ul li").length === 2 && H.app.ed.querySelectorAll("ol li").length === 2'), 'write: bulleted and numbered lists')
    check(a.js('() => H.app.ed.querySelectorAll("table td").length === 9'), 'write: the 3×3 table')
    check(a.js('() => !!H.app.ed.querySelector("img[src^=\'data:image/png\']")'), 'write: the picture')
    check(a.js('() => !!H.app.ed.querySelector("a[href=\'https://hearthos.fun\']")'), 'write: the link')
    pg = a.js('() => H.app.page')
    check(pg['size'] == 'A4' and abs(pg['margins']['l'] - 20) < 0.5 and abs(pg['margins']['t'] - 22) < 0.5, f'write: page size and margins ({pg["size"]}, {pg["margins"]})')
    check(pg.get('header') == 'Hearth header text' and pg.get('footer') == 'Hearth footer text', f'write: header and footer ({pg.get("header")!r}, {pg.get("footer")!r})')
    a.js('() => H.hideStart()')
    a.shot('docx-opened')

    # Save as .docx and read it back with python-docx.
    a.js('() => { H.app.page.numbers = "footer-center"; }')
    n = a.save_as('/docs/out.docx')
    out = os.path.join(tmp, 'out.docx')
    open(out, 'wb').write(a.get('/docs/out.docx'))
    check(n > 2000, f'write: saved .docx ({n} bytes)')
    try:
        d = Document(out)
        text = '\n'.join(p.text for p in d.paragraphs)
        for want in ['Quarterly report', 'bold words', 'Step two', 'Ünïcödé', 'A centred line']:
            check(want in text, f'write: .docx round trip keeps "{want}"')
        check(any(p.style.name.startswith('Heading 1') for p in d.paragraphs), 'write: .docx round trip keeps Heading 1')
        check(any(r.bold for p in d.paragraphs for r in p.runs if 'bold words' in r.text), 'write: .docx round trip keeps bold')
        check(any(r.italic for p in d.paragraphs for r in p.runs if 'italic' in r.text), 'write: .docx round trip keeps italic')
        check(any(r.font.color and r.font.color.rgb is not None and str(r.font.color.rgb) == 'C00000' for p in d.paragraphs for r in p.runs), 'write: .docx round trip keeps colour')
        check(any(r.font.size and r.font.size.pt == 20 and r.font.name == 'Georgia' for p in d.paragraphs for r in p.runs), 'write: .docx round trip keeps size and font')
        check(len(d.tables) == 1 and d.tables[0].cell(1, 2).text == 'R2C3', 'write: .docx round trip keeps the table')
        check(len(d.inline_shapes) == 1, 'write: .docx round trip keeps the picture')
        sec = d.sections[0]
        check(abs(sec.page_width.mm - 210) < 1 and abs(sec.left_margin.mm - 20) < 1, 'write: .docx round trip keeps page size and margins')
        check('Hearth header text' in sec.header.paragraphs[0].text, 'write: .docx round trip keeps the header')
        foot = ' '.join(p.text for p in sec.footer.paragraphs)
        check('Hearth footer text' in foot, 'write: .docx round trip keeps the footer')
        fx = zipfile.ZipFile(out).read([n for n in zipfile.ZipFile(out).namelist() if n.startswith('word/footer')][0]).decode()
        check('PAGE' in fx, 'write: .docx has page numbers in the footer')
        dx = zipfile.ZipFile(out).read('word/document.xml').decode()
        check('w:hyperlink' in dx, 'write: .docx round trip keeps the link')
        check('w:numPr' in dx, 'write: .docx round trip keeps lists')
        check(any(p.alignment is not None and int(p.alignment) == 1 for p in d.paragraphs if 'centred' in p.text), 'write: .docx round trip keeps centring')
    except Exception as e:
        check(False, f'write: python-docx reads the saved .docx ({e})')
        traceback.print_exc()

    # Open our own .docx again: same text.
    r = a.open('/docs/out.docx', a.get('/docs/out.docx'))
    html2 = a.js('() => H.app.ed.innerText')
    check(r is True and 'Quarterly report' in html2 and 'R2C3' in html2, 'write: opens its own .docx again')
    check(a.js('() => H.app.page.numbers') == 'footer-center', 'write: page numbers come back from its own .docx')

    # Every other format: save, reopen, the text is still there.
    for ext in ('odt', 'rtf', 'html', 'md', 'txt'):
        try:
            n = a.save_as(f'/docs/out.{ext}')
            r = a.open(f'/docs/out.{ext}', a.get(f'/docs/out.{ext}'))
            t = a.js('() => H.app.ed.innerText')
            check(r is True and 'Quarterly report' in t and 'Step two' in t and 'R2C3' in t, f'write: .{ext} save and reopen keep the text ({n} bytes)')
            if ext in ('odt', 'rtf', 'html'):
                check(a.js('() => !!H.app.ed.querySelector("b") && !!H.app.ed.querySelector("table") && !!H.app.ed.querySelector("li")'), f'write: .{ext} keeps bold, the table and the lists')
            if ext in ('odt', 'html', 'rtf'):
                check(a.js('() => !!H.app.ed.querySelector("img")'), f'write: .{ext} keeps the picture')
            if ext == 'md':
                check(a.js('() => !!H.app.ed.querySelector("h1") && !!H.app.ed.querySelector("table")'), 'write: .md keeps headings and the table')
            open(os.path.join(tmp, f'out.{ext}'), 'wb').write(a.get(f'/docs/out.{ext}'))
        except Exception as e:
            check(False, f'write: .{ext} ({e})')
    # LibreOffice-style check of the .odt: it's a valid zip with the right mimetype first.
    z = zipfile.ZipFile(os.path.join(tmp, 'out.odt'))
    check(z.namelist()[0] == 'mimetype' and z.read('mimetype') == b'application/vnd.oasis.opendocument.text', 'write: .odt starts with its mimetype')
    check(b'Quarterly report' in z.read('content.xml'), 'write: .odt content.xml has the text')
    rtf = open(os.path.join(tmp, 'out.rtf'), 'rb').read()
    check(rtf.startswith(b'{\\rtf1') and b'Quarterly report' in rtf, 'write: .rtf is RTF with the text')

    # Editing: type, bold, undo, find and replace, word count, a table, page setup, print layout.
    a.js('() => { H.app.newDoc("blank"); H.doc = { name: "Untitled", used: false }; H.hideStart(); H.app.focus(); }')
    a.page.keyboard.type('Hello Hearth wrold. ')
    a.page.keyboard.press('Control+B')
    a.page.keyboard.type('Bold bit')
    a.page.keyboard.press('Control+B')
    check(a.js('() => H.app.ed.querySelector("b") && H.app.ed.querySelector("b").textContent') == 'Bold bit', 'write: Ctrl+B makes bold text while typing')
    a.page.keyboard.press('Control+Z')
    check('Bold bit' not in a.js('() => H.app.ed.innerText'), 'write: Ctrl+Z undoes typing')
    a.page.keyboard.press('Control+Y')
    check('Bold bit' in a.js('() => H.app.ed.innerText'), 'write: Ctrl+Y redoes it')
    a.js('() => { document.querySelector("#fFind").value = "Hearth"; document.querySelector("#fRepl").value = "Ember"; H.app.replaceAll(); }')
    check('Hello Ember' in a.js('() => H.app.ed.innerText'), 'write: replace all')
    check(a.js('() => H.app.counts().words') == 5, f'write: word count ({a.js("() => H.app.counts().words")})')
    a.js('() => H.app.insertTable(2, 3)')
    check(a.js('() => H.app.ed.querySelectorAll("table td").length') == 6, 'write: inserts a 2×3 table')
    a.page.wait_for_timeout(2500)
    spell = a.js('() => H.app.spell.ready && H.app.spell.isWrong("wrold") && !H.app.spell.isWrong("Hello")')
    check(spell, 'write: the spell checker knows "wrold" is wrong and "Hello" is right')
    sug = a.js('() => H.app.spell.suggest("wrold")')
    check('world' in (sug or []), f'write: suggests "world" ({sug})')
    a.shot('editing')
    # Long document: pages, header, footer, numbers.
    a.js('() => { H.app.setDoc(Array.from({length: 120}, (_, i) => `<p>Paragraph ${i + 1}: ` + "words ".repeat(30) + "</p>").join(""), { header: "Head", footer: "Foot", numbers: "footer-center", size: "A4" }); }')
    pages = a.js('() => H.app.paginate(null).length')
    check(pages >= 8, f'write: a long document is laid out on pages ({pages})')
    a.js('() => H.app.togglePreview()')
    a.page.wait_for_timeout(300)
    check(a.js('() => document.querySelectorAll("#preview .ppage").length') == pages, 'write: print layout shows every page')
    check('Page 2 of' in a.js('() => document.querySelector("#preview").innerText'), 'write: page numbers in the print layout')
    a.shot('print-layout')
    a.js('() => H.app.togglePreview()')
    # PDF through the browser's own PDF printer (Chromium only), to check the print CSS gives one PDF page per page.
    if browser.browser_type.name == 'chromium':
        a.js('() => H.preparePrint()')
        pdf = a.page.pdf(prefer_css_page_size=True)
        npdf = pdf.count(b'/Type /Page') - pdf.count(b'/Type /Pages')
        check(abs(npdf - pages) <= 1, f'write: the PDF has a page per page ({npdf} for {pages})')
        open(os.path.join(tmp, 'write.pdf'), 'wb').write(pdf)
    # The unsaved-changes guard.
    a.js('() => H.changed()')
    a.js('() => { H.mock.emit("office://close"); }')
    a.page.wait_for_selector('.dlg')
    check('Save changes' in a.js('() => document.querySelector(".dlg h3").textContent'), 'write: closing with changes asks to save')
    a.shot('close-guard')
    a.page.click('.dlg button[data-v="discard"]')
    check(a.js('() => H.mock.quit'), 'write: "Don\'t save" closes')
    check(not a.errors, f'write: no script errors ({a.errors[:3]})')
    a.page.close()


# ------------------------------------------------------------------ Sheets
def test_sheets(browser, base, files, shots, tmp):
    import csv as csvmod
    from openpyxl import load_workbook
    a = App(browser, base, 'sheets', shots)
    a.shot('start')
    r = a.open('/docs/sample.xlsx', open(files['sample.xlsx'], 'rb').read())
    check(r is True, f'sheets: opens an .xlsx made by openpyxl ({r})')

    def v(ref, sheet=None):
        return a.js('([ref, sh]) => { const S = H.app; const p = S.parseAddr(ref); const s = sh ? S.sheets.find(x => x.name === sh) : S.sh(); return S.val(p.r, p.c, s.id); }', [ref, sheet])

    def t(ref):
        return a.js('(ref) => { const S = H.app; const p = S.parseAddr(ref); return S.text(p.r, p.c); }', ref)

    def ser(r, c):
        return a.js('([r, c]) => H.app.ser(r, c)', [r, c])

    def fmt(r, c):
        return a.js('([r, c]) => H.app.fmtAt(r, c)', [r, c]) or {}

    check(a.js('() => H.app.sheets.map(s => s.name).join(",")') == 'Budget,Notes', 'sheets: both sheets are there')
    check(v('D2') == 10, f'sheets: =B2*C2 is 10 ({v("D2")})')
    check(abs((v('D7') or 0) - 29.25) < 1e-9, f'sheets: =SUM is 29.25 ({v("D7")})')
    check(v('D8') == 2.84, f'sheets: =ROUND(AVERAGE()) is 2.84 ({v("D8")})')
    check(v('D9') == 1.2, f'sheets: =VLOOKUP(...,FALSE) is 1.2 ({v("D9")})')
    check(v('D10') == 'over', f'sheets: =IF is "over" ({v("D10")})')
    check(v('A2', 'Notes') == 58.5, f'sheets: a formula across sheets ({v("A2", "Notes")})')
    f = fmt(0, 0)
    check(f.get('b') and f.get('bg') == '#2f5496' and f.get('color') == '#ffffff', f'sheets: bold, fill and font colour read ({f})')
    check(t('B2') == '$2.50', f'sheets: currency format shows $2.50 ({t("B2")})')
    check(t('D7') == '29.25', f'sheets: 0.00 format ({t("D7")})')
    check(fmt(2, 0).get('align') == 'right', 'sheets: alignment read')
    check((fmt(4, 0).get('bd') or {}).get('t') == 'thin', 'sheets: borders read')
    m = a.js('() => H.app.meta()')
    check(m['freeze'] == {'r': 1, 'c': 0}, f'sheets: frozen top row read ({m["freeze"]})')
    check({'r': 0, 'c': 5, 'r2': 0, 'c2': 6} in m['merges'], 'sheets: merged cells read')
    check(120 < m['colW'].get('0', 0) < 150, f'sheets: column width read ({m["colW"].get("0")})')
    a.js('() => H.hideStart()')
    a.page.wait_for_timeout(300)
    a.shot('xlsx-opened')

    # Typing: a formula into a cell, through the keyboard.
    a.js('() => { H.app.select(11, 0); H.app.focus(); }')
    a.page.keyboard.type('=sum(b2:b5')
    a.page.keyboard.press('Enter')
    check(abs((v('A12') or 0) - 11.35) < 1e-9, f'sheets: typed =sum(b2:b5 gives 11.35 ({v("A12")}, {ser(11, 0)})')
    a.page.keyboard.type('1,234.5')
    a.page.keyboard.press('Enter')
    check(v('A13') == 1234.5 and t('A13') == '1,234.5', f'sheets: 1,234.5 is a number ({v("A13")}, {t("A13")})')
    a.page.keyboard.type('2026-03-04')
    a.page.keyboard.press('Enter')
    check(t('A14') == '2026-03-04' and isinstance(v('A14'), (int, float)), f'sheets: a date is a date ({v("A14")}, {t("A14")})')
    a.page.keyboard.press('Control+Z')
    check(v('A14') is None, 'sheets: Ctrl+Z undoes the last entry')
    a.page.keyboard.press('Control+Y')
    check(t('A14') == '2026-03-04', 'sheets: Ctrl+Y redoes it')
    # Fill: 1, 2 → 3, 4, 5; a formula fills with its references moving; Mon → Wed.
    a.js('() => { const S = H.app; S.hf.setCellContents({sheet: S.sid(), row: 19, col: 0}, [[1], [2]]); S.fill({r: 19, c: 0, r2: 20, c2: 0}, {r: 19, c: 0, r2: 23, c2: 0}); }')
    got = [v(f'A{i}') for i in range(20, 25)]
    check(got == [1, 2, 3, 4, 5], f'sheets: fill continues 1, 2 → 5 ({got})')
    a.js('() => { const S = H.app; S.hf.setCellContents({sheet: S.sid(), row: 19, col: 1}, [["=A20*10"]]); S.fill({r: 19, c: 1, r2: 19, c2: 1}, {r: 19, c: 1, r2: 23, c2: 1}); }')
    check(v('B24') == 50 and ser(23, 1) == '=A24*10', f'sheets: a filled formula moves its references ({v("B24")}, {ser(23, 1)})')
    a.js('() => { const S = H.app; S.hf.setCellContents({sheet: S.sid(), row: 19, col: 2}, [["Mon"]]); S.fill({r: 19, c: 2, r2: 19, c2: 2}, {r: 19, c: 2, r2: 21, c2: 2}); }')
    check(v('C22') == 'Wed', f'sheets: Mon fills to Wed ({v("C22")})')
    # Sort the item table by name: totals (formulas) follow their rows.
    a.js('() => { H.app.select(1, 0); H.app.sort({r: 0, c: 0, r2: 4, c2: 3}, [{c: 0, dir: 1}], true); }')
    got = [v(f'A{i}') for i in range(2, 6)]
    check(got == ['Apples', 'Bread', 'Cheese', 'Milk'], f'sheets: sort A→Z ({got})')
    check(abs((v('D2') or 0) - 4.8) < 1e-9 and ser(1, 3) == '=B2*C2', f'sheets: formulas move with their rows when sorting ({v("D2")}, {ser(1, 3)})')
    check(abs((v('D7') or 0) - 29.25) < 1e-9, 'sheets: the total is the same after sorting')
    # Filter: hide Milk.
    a.js('() => { const S = H.app; S.meta().filter = {r: 0, c: 0, r2: 4, c2: 3, crit: {0: ["Milk"]}}; S.applyFilter(); S.layout(); S.draw(); }')
    check(a.js('() => !!H.app.meta().hiddenRows[4] && !H.app.meta().hiddenRows[1]'), 'sheets: a filter hides the rows it should')
    a.page.wait_for_timeout(200)
    a.shot('sorted-filtered')
    a.js('() => { const S = H.app; S.meta().filter = null; S.applyFilter(); S.layout(); S.draw(); }')
    # Insert a row above row 2: references below move; undo puts it back.
    a.js('() => { H.app.select(1, 0); H.app.insert("row", true); }')
    check(ser(7, 3) == '=SUM(D3:D6)', f'sheets: inserting a row moves references ({ser(7, 3)})')
    a.js('() => H.app.undo()')
    check(ser(6, 3) == '=SUM(D2:D5)', f'sheets: undo puts the row back ({ser(6, 3)})')
    # Formatting from the toolbar.
    a.js('() => { H.app.select(0, 1, 0, 3); H.app.setFmt({ b: true, bg: "#ffff00" }); H.app.borders("outer"); }')
    check(fmt(0, 3).get('bg') == '#ffff00', 'sheets: fill colour set')

    # Save as .xlsx; openpyxl reads it.
    n = a.save_as('/docs/out.xlsx')
    out = os.path.join(tmp, 'out.xlsx')
    open(out, 'wb').write(a.get('/docs/out.xlsx'))
    check(n > 3000, f'sheets: saved .xlsx ({n} bytes)')
    try:
        wb = load_workbook(out)
        ws = wb['Budget']
        check(wb.sheetnames == ['Budget', 'Notes'], f'sheets: .xlsx keeps both sheets ({wb.sheetnames})')
        check(ws['D7'].value == '=SUM(D2:D5)', f'sheets: .xlsx keeps formulas ({ws["D7"].value})')
        check(str(ws['D9'].value).upper().startswith('=VLOOKUP("MILK"') and str(ws['D9'].value).endswith('FALSE)'), f'sheets: VLOOKUP with FALSE saved the Excel way ({ws["D9"].value})')
        check(wb['Notes']['A2'].value == '=Budget!D7*2', f'sheets: .xlsx keeps cross-sheet formulas ({wb["Notes"]["A2"].value})')
        check(ws['A1'].font.b and str(ws['A1'].fill.fgColor.rgb).endswith('2F5496'), f'sheets: .xlsx keeps bold and fill ({ws["A1"].font.b}, {ws["A1"].fill.fgColor.rgb})')
        check(ws['A1'].font.color is not None and str(ws['A1'].font.color.rgb).endswith('FFFFFF'), 'sheets: .xlsx keeps font colour')
        check('$' in ws['B3'].number_format or '$' in ws['B2'].number_format, f'sheets: .xlsx keeps the currency format ({ws["B2"].number_format}, {ws["B3"].number_format})')
        check(ws['D7'].number_format == '0.00', f'sheets: .xlsx keeps 0.00 ({ws["D7"].number_format})')
        check(ws.freeze_panes == 'A2', f'sheets: .xlsx keeps the frozen row ({ws.freeze_panes})')
        check('F1:G1' in [str(x) for x in ws.merged_cells.ranges], 'sheets: .xlsx keeps merged cells')
        check(ws['B1'].border.top.style == 'thin' and str(ws['D1'].fill.fgColor.rgb).endswith('FFFF00'), 'sheets: .xlsx keeps borders and new fills')
        check(ws['A2'].border.left.style == 'thin', 'sheets: .xlsx keeps borders it read (moved with its row by the sort)')
        check(ws['A5'].alignment.horizontal == 'right', 'sheets: .xlsx keeps alignment (moved with its row by the sort)')
        check(abs(ws.column_dimensions['A'].width - 18) < 2, f'sheets: .xlsx keeps column width ({ws.column_dimensions["A"].width})')
        wv = load_workbook(out, data_only=True)['Budget']
        check(abs((wv['D7'].value or 0) - 29.25) < 1e-9, f'sheets: .xlsx has the formulas\' results for other programs ({wv["D7"].value})')
    except Exception as e:
        check(False, f'sheets: openpyxl reads the saved .xlsx ({e})')
        traceback.print_exc()
    # Its own .xlsx back in.
    r = a.open('/docs/out.xlsx', a.get('/docs/out.xlsx'))
    check(r is True and abs((v('D7') or 0) - 29.25) < 1e-9 and fmt(0, 0).get('b'), 'sheets: opens its own .xlsx again with formatting')
    # .ods: values and formulas survive.
    n = a.save_as('/docs/out.ods')
    r = a.open('/docs/out.ods', a.get('/docs/out.ods'))
    check(r is True and abs((v('D7') or 0) - 29.25) < 1e-9 and str(ser(6, 3)).upper().startswith('=SUM'), f'sheets: .ods save and reopen keep values and formulas ({n} bytes, {ser(6, 3)})')
    check(a.js('() => H.app.sheets.length') == 2, 'sheets: .ods keeps both sheets')
    # CSV.
    r = a.open('/docs/sample.csv', open(files['sample.csv'], 'rb').read())
    check(r is True and v('A2') == 'Smith, Jo' and v('C3') == 'New\nYork' and v('A4') == 'Zoë' and v('B3') == 78.5, f'sheets: CSV with quotes, commas, line breaks and accents ({v("A2")!r}, {v("C3")!r}, {v("B3")!r})')
    a.save_as('/docs/out.csv')
    rows = list(csvmod.reader(io.StringIO(a.get('/docs/out.csv').decode('utf-8'))))
    check(rows[1] == ['Smith, Jo', '91', 'Paris'] and rows[2][2] == 'New\nYork' and rows[3][2] == 'Kraków', f'sheets: CSV round trip ({rows[:4]})')
    # A new spreadsheet from a template works out its formulas.
    a.js('() => { H.app.newDoc("invoice"); }')
    check(abs((v('D12') or 0) - 99) < 1e-9, f'sheets: the invoice template adds up ({v("D12")})')
    a.page.wait_for_timeout(200)
    a.shot('template')
    check(not a.errors, f'sheets: no script errors ({a.errors[:3]})')
    a.page.close()


# ------------------------------------------------------------------ Slides
def test_slides(browser, base, files, shots, tmp):
    from pptx import Presentation
    a = App(browser, base, 'slides', shots)
    a.shot('start')
    r = a.open('/docs/sample.pptx', open(files['sample.pptx'], 'rb').read())
    check(r is True, f'slides: opens a .pptx made by python-pptx ({r})')
    P = a.js('() => H.app.P')
    check(len(P['slides']) == 3 and P['h'] == 540, f'slides: 3 slides, 16:9 ({len(P["slides"])}, {P["h"]})')
    s1 = P['slides'][0]
    texts1 = ' '.join(e.get('html', '') for e in s1['els'])
    check('Hearth Slides test' in texts1 and 'A subtitle line' in texts1, 'slides: title and subtitle text')
    check(any(e.get('role') == 'title' and 'Hearth Slides test' in e.get('html', '') for e in s1['els']), 'slides: the title placeholder is a title')
    check(s1['notes'] == 'Speaker notes for slide one', f'slides: speaker notes ({s1["notes"]!r})')
    s2 = P['slides'][1]
    body = next((e for e in s2['els'] if 'First point' in e.get('html', '')), None)
    check(body is not None and body['html'].count('<li') == 3, f'slides: the bullet list ({body and body["html"]})')
    s3 = P['slides'][2]
    big = next((e for e in s3['els'] if 'Big red words' in e.get('html', '')), None)
    check(big is not None and ('c00000' in big['html'].lower() or (big.get('color') or '').lower() == '#c00000') and '<b>' in big['html'], f'slides: colour and bold kept ({big and big["html"]})')
    check(big is not None and ('53.3px' in big['html'] or abs(big.get('size', 0) - 53) <= 1), f'slides: 40 pt text size kept ({big and (big.get("size"), big["html"])})')
    box = next((e for e in s3['els'] if e['type'] == 'shape'), None)
    check(box is not None and box.get('fill') == '#1f9a55' and 'In a box' in box.get('html', ''), f'slides: a filled shape with text ({box})')
    check(any(e['type'] == 'image' and e['src'].startswith('data:image/png') for e in s3['els']), 'slides: the picture')
    a.js('() => H.hideStart()')
    a.page.wait_for_timeout(300)
    a.js('() => H.app.go(2)')
    a.page.wait_for_timeout(300)
    a.shot('pptx-opened')

    # Save as .pptx: python-pptx reads it back.
    n = a.save_as('/docs/out.pptx')
    out = os.path.join(tmp, 'out.pptx')
    open(out, 'wb').write(a.get('/docs/out.pptx'))
    check(n > 5000, f'slides: saved .pptx ({n} bytes)')
    try:
        p = Presentation(out)
        check(len(p.slides) == 3, f'slides: .pptx has 3 slides ({len(p.slides)})')
        alltext = [' '.join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in p.slides]
        check('Hearth Slides test' in alltext[0] and 'Third point' in alltext[1] and 'Big red words' in alltext[2], f'slides: .pptx keeps the text ({alltext})')
        check(p.slides[0].has_notes_slide and 'Speaker notes for slide one' in p.slides[0].notes_slide.notes_text_frame.text, 'slides: .pptx keeps speaker notes')
        check(sum(1 for sh in p.slides[2].shapes if sh.shape_type == 13) == 1, 'slides: .pptx keeps the picture')
        runs = [r for sh in p.slides[2].shapes if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs if 'Big red' in r.text]
        check(runs and runs[0].font.bold and str(runs[0].font.color.rgb) == 'C00000', 'slides: .pptx keeps bold and colour')
        check(runs and abs(runs[0].font.size.pt - 40) < 1, f'slides: .pptx keeps the text size ({runs and runs[0].font.size and runs[0].font.size.pt})')
        bullets = [para for sh in p.slides[1].shapes if sh.has_text_frame for para in sh.text_frame.paragraphs if 'point' in para.text]
        check(len(bullets) == 3, f'slides: .pptx keeps the three bullet points ({len(bullets)})')
        check(abs(p.slide_width / p.slide_height - 16 / 9) < 0.01, 'slides: .pptx is 16:9')
    except Exception as e:
        check(False, f'slides: python-pptx reads the saved .pptx ({e})')
        traceback.print_exc()
    r = a.open('/docs/out.pptx', a.get('/docs/out.pptx'))
    P2 = a.js('() => H.app.P')
    check(r is True and len(P2['slides']) == 3 and 'Big red words' in ' '.join(e.get('html', '') for e in P2['slides'][2]['els']), 'slides: opens its own .pptx again')
    # Hearth's own format keeps everything.
    a.save_as('/docs/out.hslides')
    before = a.js('() => JSON.stringify(H.app.P.slides.map(s => s.els.map(e => [e.type, e.x, e.y, e.w, e.h, e.html || ""])))')
    r = a.open('/docs/out.hslides', a.get('/docs/out.hslides'))
    after = a.js('() => JSON.stringify(H.app.P.slides.map(s => s.els.map(e => [e.type, e.x, e.y, e.w, e.h, e.html || ""])))')
    check(r is True and before == after, 'slides: .hslides save and reopen keep every box exactly')

    # Editing: a new deck, a title, a text box, a shape; undo.
    a.js('() => { H.app.newDoc("ember"); H.doc = { name: "Untitled", used: false }; H.hideStart(); }')
    a.js('() => H.app.startEdit(H.app.P.slides[0].els[0].id)')
    a.page.keyboard.type('My talk')
    a.js('() => H.app.endEdit()')
    check('My talk' in a.js('() => H.app.P.slides[0].els[0].html'), 'slides: typing into the title')
    a.js('() => H.app.addSlide("content")')
    a.js('() => H.app.addShape("ellipse")')
    check(a.js('() => H.app.P.slides.length') == 2 and a.js('() => H.app.P.slides[1].els.some(e => e.shape === "ellipse")'), 'slides: a new slide with a shape')
    a.js('() => H.app.undo()')
    check(not a.js('() => H.app.P.slides[1].els.some(e => e.shape === "ellipse")'), 'slides: undo removes the shape')
    a.js('() => H.app.moveSlide(1, 0)')
    check('My talk' in ' '.join(e.get('html', '') for e in a.js('() => H.app.P.slides[1].els')), 'slides: reordering slides')
    a.js('() => H.app.go(1)')
    a.page.wait_for_timeout(300)
    a.shot('editing')
    # Presenting: keys move through the slides, Escape ends.
    a.js('() => H.app.present(0)')
    a.page.wait_for_timeout(200)
    check(a.js('() => !document.querySelector("#present").hidden'), 'slides: presenting fills the window')
    a.page.keyboard.press('ArrowRight')
    check(a.js('() => H.app.show && H.app.show.i') == 1, 'slides: → goes to the next slide')
    a.page.keyboard.press('s')
    a.page.wait_for_timeout(200)
    a.shot('presenter-view')
    check(a.js('() => !!document.querySelector(".pv-notes")'), 'slides: S shows the presenter view with notes')
    a.js('() => H.padPress("LB")')
    check(a.js('() => H.app.show && H.app.show.i') == 0, 'slides: LB on a controller goes back')
    a.page.keyboard.press('Escape')
    check(a.js('() => document.querySelector("#present").hidden'), 'slides: Escape stops presenting')
    # Print layout: a page per slide.
    pages = a.js('async () => { const d = await H.preparePrint(); return [document.querySelectorAll("#print .ppage").length, d.w, d.h]; }')
    check(pages[0] == 2 and abs(pages[1] - 254) < 0.1, f'slides: print layout is a 10-inch page per slide ({pages})')
    if browser.browser_type.name == 'chromium':
        pdf = a.page.pdf(prefer_css_page_size=True)
        npdf = pdf.count(b'/Type /Page') - pdf.count(b'/Type /Pages')
        check(npdf == 2, f'slides: the PDF has a page per slide ({npdf})')
    check(not a.errors, f'slides: no script errors ({a.errors[:3]})')
    a.page.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--browser', default='chromium')
    ap.add_argument('--shots', default=None)
    ap.add_argument('--only', default='write,sheets,slides')
    args = ap.parse_args()
    from playwright.sync_api import sync_playwright
    tmp = tempfile.mkdtemp(prefix='hearth-office-')
    files = samples.make_all(os.path.join(tmp, 'samples'))
    if args.shots:
        os.makedirs(args.shots, exist_ok=True)
    srv = serve()
    base = f'http://127.0.0.1:{srv.server_address[1]}'
    with sync_playwright() as p:
        browser = getattr(p, args.browser).launch()
        for name in args.only.split(','):
            fn = globals().get('test_' + name)
            if not fn:
                continue
            try:
                fn(browser, base, files, args.shots, tmp)
            except Exception as e:
                check(False, f'{name}: {e}')
                traceback.print_exc()
        browser.close()
    print(f'\n{len(PASSES)} passed, {len(FAILS)} failed ({args.browser})')
    if FAILS:
        print('Failed:\n  ' + '\n  '.join(FAILS))
    sys.exit(1 if FAILS else 0)


if __name__ == '__main__':
    main()
